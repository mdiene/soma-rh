import openpyxl, json, datetime, sys
xls = sys.argv[1]; out_ref = sys.argv[2]; out_data = sys.argv[3]
wv = openpyxl.load_workbook(xls, data_only=True)

def q(v):
    if v is None: return 'null'
    if isinstance(v, bool): return 'true' if v else 'false'
    if isinstance(v, (int, float)): return repr(v)
    if isinstance(v, datetime.datetime): return "'" + v.date().isoformat() + "'"
    if isinstance(v, datetime.date): return "'" + v.isoformat() + "'"
    s = str(v).strip()
    if s == '': return 'null'
    return "'" + s.replace("'", "''") + "'"

ref = []
ref.append("-- SOMA-RH — 05 : données de référence issues du classeur Excel (PARAMETRES)\n")
ref.append("insert into public.payroll_settings (effective_from) values ('2023-01-01');\n")
ref.append("""
with s as (insert into public.tax_slabs (name, effective_from)
           values ('Barème IR Sénégal (CGI 2013) — Excel SOMA', '2013-01-01') returning id)
, b as (insert into public.tax_slab_brackets (slab_id, from_amount, rate)
        select s.id, x.f, x.r from s, (values (0,0),(630000,0.20),(1500000,0.30),(4000000,0.35),(8000000,0.37),(13500000,0.40)) as x(f,r))
, f as (insert into public.ir_family_reductions (slab_id, parts, rate, min_amount, max_amount)
        select s.id, x.p, x.r, x.mn, x.mx from s, (values
          (1.0,0,0,0),(1.5,0.10,100000,300000),(2.0,0.15,200000,650000),(2.5,0.20,300000,1100000),
          (3.0,0.25,400000,1650000),(3.5,0.30,500000,2030000),(4.0,0.35,600000,2490000),
          (4.5,0.40,700000,2755000),(5.0,0.45,800000,3180000)) as x(p,r,mn,mx))
insert into public.trimf_brackets (slab_id, min_annual_gross, annual_amount, per_part)
select s.id, x.m, x.a, x.pp from s, (values
  (0,900,false),(600000,3600,true),(1000000,4800,true),(2000000,12000,true),(7000000,18000,true),(12000000,36000,true)) as x(m,a,pp);
""")
comps = [
 ('SAL_BASE','Salaire de base',10,'gross_earning','fixed'),
 ('SURSALAIRE','Sursalaire',11,'gross_earning','fixed'),
 ('PRIME_ANC',"Prime d'ancienneté",12,'gross_earning','computed'),
 ('HONORAIRES','Honoraires prestataires',13,'gross_earning','fixed'),
 ('AUTRE_FIXE','Autre élément fixe',14,'gross_earning','fixed'),
 ('RET_ABSENCE','Retenues absences',15,'gross_earning','computed'),
 ('PRIME_PERF','Prime de performance',16,'gross_earning','variable'),
 ('GRATIF_TABASKI','Gratification Tabaski',17,'gross_earning','variable'),
 ('FORFAIT_HS','Forfait heures sup',18,'gross_earning','variable'),
 ('PRIME_DIVERSE','Prime diverse',19,'gross_earning','variable'),
 ('IND_MALADIE','Indemnité de maladie',22,'gross_earning','computed'),
 ('ALLOC_CONGE','Allocation congé',23,'gross_earning','computed'),
 ('IND_PRECARITE','Indemnité de fin de contrat (précarité)',24,'gross_earning','computed'),
 ('HS_15','Heures sup. majorées à 15%',25,'gross_earning','computed'),
 ('HS_40','Heures sup. majorées à 40%',26,'gross_earning','computed'),
 ('HS_60','Heures sup. majorées à 60%',27,'gross_earning','computed'),
 ('HS_100','Heures sup. majorées à 100%',28,'gross_earning','computed'),
 ('IR','Impôt sur le revenu',51,'statutory_deduction','computed'),
 ('TRIMF','Retenue TRIMF',52,'statutory_deduction','computed'),
 ('IPRES_RG','IPRES Régime Général',53,'statutory_deduction','computed'),
 ('IPRES_RC','IPRES Régime Complémentaire Cadre',54,'statutory_deduction','computed'),
 ('IPM','Cotisations IPM',55,'statutory_deduction','computed'),
 ('BRS','BRS (5%) prestataires',56,'statutory_deduction','computed'),
 ('CSS_AF','CSS - Allocations familiales',57,'employer_contribution','computed'),
 ('CSS_AT','CSS - Accident de travail',58,'employer_contribution','computed'),
 ('CFCE','CFCE',59,'employer_contribution','computed'),
 ('IND_TRANSPORT','Indemnité de transport',70,'non_taxable_earning','computed'),
 ('IND_PANIER','Indemnité de panier',71,'non_taxable_earning','computed'),
 ('IND_KM','Indemnité kilométrique',72,'non_taxable_earning','computed'),
 ('IND_LICENCIEMENT','Indemnité de licenciement',73,'non_taxable_earning','computed'),
 ('IND_RETRAITE','Indemnité de départ à la retraite',74,'non_taxable_earning','computed'),
 ('RET_AVANCE_CAISSE','Retenue avance caisse',75,'net_deduction','variable'),
 ('RET_COOPERATIVE','Retenue coopérative',76,'net_deduction','variable'),
 ('RET_AVANCE_FIN','Retenue avance financière',77,'net_deduction','variable'),
 ('RET_FRAIS_MED','Frais médicaux',78,'net_deduction','variable'),
 ('RET_CAISSE_COMMUNE','Caisse commune',79,'net_deduction','variable'),
 ('RET_AVANCE_KORITE','Retenue avance Korité',80,'net_deduction','variable'),
 ('APPOINT_PREC','Appoint du mois précédent',81,'rounding','computed'),
 ('APPOINT_COURANT','Appoint du mois en cours',82,'rounding','computed'),
]
ref.append("insert into public.salary_components (code, name, rubric_number, component_type, input_mode) values\n" +
           ",\n".join(f"  ({q(c)}, {q(n)}, {r}, {q(t)}, {q(m)})" for c,n,r,t,m in comps) + ";\n")
ref.append("""insert into public.leave_types (code, name, is_paid, affects_payroll, days_per_month, max_days_per_year) values
  ('CONGE_ANNUEL','Congé annuel payé', true, 'none', 2, 24),
  ('MALADIE','Congé maladie', true, 'sick_leave', 0, null),
  ('ABSENCE_NP','Absence non payée', false, 'absence', 0, null),
  ('PERMISSION','Permission exceptionnelle', true, 'none', 0, 10),
  ('MATERNITE','Congé de maternité', true, 'none', 0, 98);
""")
ws = wv['PARAMETRES']
agreements = []
for r in ws.iter_rows(min_row=31, max_row=41, min_col=11, max_col=11, values_only=True):
    if r[0]: agreements.append(r[0])
ref.append("insert into public.collective_agreements (code, name) values\n" + ",\n".join(
    f"  ({q(a)}, {q(a.replace('_',' ').title())})" for a in agreements) + ";\n")
grades = []
for r in ws.iter_rows(min_row=15, max_row=268, max_col=6, values_only=True):
    if r[0] is None: continue
    grades.append(f"  ({q(r[4])}, {q(' '.join(str(r[0]).split()))}, {r[1]!r}, {r[2]!r}, {q(r[3])}, {q(r[5])})")
ref.append("insert into public.employee_grades (agreement_code, code, hourly_rate, base_salary, staff_status, grade_class) values\n"
           + ",\n".join(grades) + "\non conflict do nothing;\n")
ref.append("""insert into public.companies (name, legal_form, share_capital, address, phone)
values ('SOCIETE MINIERE AFRICAINE', 'SARL', 10000000, '11 rue Alfred Goux, Apt N1, 1er étage, Dakar', '+221 77 247 25 00');

insert into public.salary_structures (company_id, name, description)
select id, 'Régime général Sénégal', 'Structure unique reprenant toutes les rubriques du classeur Excel de paie SOMA (août 2026).'
from public.companies where name = 'SOCIETE MINIERE AFRICAINE';

insert into public.salary_structure_components (structure_id, component_code, sort_order)
select s.id, c.code, c.rubric_number from public.salary_structures s cross join public.salary_components c
where s.name = 'Régime général Sénégal';
""")
open(out_ref, 'w').write("\n".join(ref))

# ------------------------------------------------------------ données salariés
fp = wv['FICHE DE PERSONNEL']
hdr = [c.value for c in fp[1]]
rows = [dict(zip(hdr, r)) for r in fp.iter_rows(min_row=2, max_row=27, values_only=True) if r[0]]
el = wv['élements de salaires']
eh = [str(c.value).replace('\n',' ').strip() for c in el[1]]
erows = {r[0]: dict(zip(eh, r)) for r in el.iter_rows(min_row=2, max_row=26, values_only=True) if r[0]}

def norm(s): return ' '.join(str(s).split()) if s is not None else None
def phone(v):
    if v is None or str(v).strip()=='' : return None
    return str(int(v)) if isinstance(v,(int,float)) else str(v).strip()

d = ["-- SOMA-RH — données salariés (FICHE DE PERSONNEL + élements de salaires, août 2026)\n"]
designs = sorted({norm(r['fonction']) for r in rows if r['fonction']})
d.append("insert into public.designations (name) values " + ", ".join(f"({q(x)})" for x in designs) + " on conflict do nothing;\n")
vals = []
for r in rows:
    mat = r['MATRICULE']
    vals.append("(" + ", ".join([
        q(mat), q(norm(r['PRENOM'])), q(norm(r['NOM'])), q(r['SEXE']), q(r['date de Naissance']),
        q(r['LIEU DE NAISSANCE']), q(r['nationalité']), q(r['PÈRE']), q(r['MERE']),
        q(phone(r['CNI ou PASSEPORT'])), q(r['DATE DE  DELIVRANCE'] if isinstance(r['DATE DE  DELIVRANCE'], datetime.datetime) else None),
        q(r["DATE DE  D'EXPIRATION"]), q(r['NINEA']), q(r["NUMERO D'AFFILIATION IPRES"]), q(r['N° ALLOCATION FAMILIALE']),
        q(r['adresse']), q(phone(r['numéro de téléphone perso'])), q(phone(r['numéro de téléphone pro'])),
        q(r['e-mail perso']), q(r['e-mail pro']), q(r["personne à contacter en cas d'urgence"]), q(phone(r['Téléphone'])),
        q(r['Situation de famille']), str(r['FEMMES'] or 0), str(r['ENFANTS'] or 0), repr(r['PART TRIMF'] or 1), repr(r['QUOTIENT FAMILIAL'] or 1),
        q(norm(r['fonction'])), q(r['CONVENTION COLLECTIVE']), q(norm(r['CATEGORIE'])), q(r['STATUT']),
        q(r['Types de contrat']), q(r["date d'entrée"]), q(r['DATE DEBUT']), q(r['DATE DE FIN']),
        q(r['date sortie']), q(r['motif']), q(r['Date RETOUR DE CONGE'])]) + ")")
d.append("""with src(employee_number, first_name, last_name, gender, date_of_birth, place_of_birth, nationality, father_name, mother_name,
  id_document_number, id_issue_date, id_expiry_date, ninea, ipres_number, css_number, address, phone_personal, phone_work,
  email_personal, email_work, emergency_contact_name, emergency_contact_phone, marital_status, spouses_count, children_count,
  trimf_parts, ir_parts, designation, agreement_code, grade_code, staff_status, employment_type, date_of_joining,
  contract_start, contract_end, relieving_date, exit_reason, leave_return_date) as (values
""" + ",\n".join(vals) + """
)
insert into public.employees (company_id, employee_number, first_name, last_name, gender, date_of_birth, place_of_birth,
  nationality, father_name, mother_name, id_document_number, id_issue_date, id_expiry_date, ninea, ipres_number, css_number,
  address, phone_personal, phone_work, email_personal, email_work, emergency_contact_name, emergency_contact_phone,
  marital_status, spouses_count, children_count, trimf_parts, ir_parts, designation_id, agreement_code, grade_id,
  employment_type, date_of_joining, contract_start, contract_end, relieving_date, exit_reason, leave_return_date)
select c.id, s.employee_number, s.first_name, s.last_name, s.gender, s.date_of_birth::date, s.place_of_birth,
  s.nationality, s.father_name, s.mother_name, s.id_document_number, s.id_issue_date::date, s.id_expiry_date::date,
  s.ninea, s.ipres_number, s.css_number, s.address, s.phone_personal, s.phone_work, s.email_personal, s.email_work,
  s.emergency_contact_name, s.emergency_contact_phone, s.marital_status, s.spouses_count::smallint, s.children_count::smallint,
  s.trimf_parts::numeric, s.ir_parts::numeric, d.id, s.agreement_code,
  (select g.id from public.employee_grades g where g.agreement_code = s.agreement_code and g.code = s.grade_code
     and lower(g.staff_status) = lower(s.staff_status) order by g.effective_from desc limit 1),
  s.employment_type, s.date_of_joining::date, s.contract_start::date, s.contract_end::date,
  s.relieving_date::date, s.exit_reason, s.leave_return_date::date
from src s
cross join (select id from public.companies where name = 'SOCIETE MINIERE AFRICAINE') c
left join public.designations d on d.name = s.designation;
""")
# affectations de structure (sursalaire Excel, transport, congé)
a = []
for mat, e in erows.items():
    a.append(f"({q(mat)}, {e['SURSALAIRE'] or 0!r}, {e['honoraires prestataires'] or 0!r}, {e['xxxxxxxx'] or 0!r}, "
             f"{q(e['transport'] or 'Aucun')}, {'true' if str(e['le congé est il à payer?']).upper()=='OUI' else 'false'}, "
             f"{q(e['Net à simuler'])})")
d.append("""insert into public.salary_structure_assignments (employee_id, structure_id, from_date, sursalaire, fees, other_fixed,
  transport_mode, leave_paid_monthly, target_net)
select e.id, s.id, '2026-01-01', v.sursalaire, v.fees, v.other_fixed, v.transport_mode, v.leave_paid, v.target_net::numeric
from (values
""" + ",\n".join(a) + """
) as v(employee_number, sursalaire, fees, other_fixed, transport_mode, leave_paid, target_net)
join public.employees e on e.employee_number = v.employee_number
cross join (select id from public.salary_structures where name = 'Régime général Sénégal') s;
""")
# variables du mois d'août 2026
pi = []; adds = []
addmap = {'Prime de performance':'PRIME_PERF','gratification tabaski':'GRATIF_TABASKI','Forfait Heures Sup':'FORFAIT_HS',
          'xxx3':'PRIME_DIVERSE','Retenues avance Korité':'RET_AVANCE_KORITE','retenues avance caisse':'RET_AVANCE_CAISSE',
          'RETENUES avance financiere':'RET_AVANCE_FIN','retenues  coopérative':'RET_COOPERATIVE','FRAIS MEDICAUX':'RET_FRAIS_MED',
          'caisse commune':'RET_CAISSE_COMMUNE'}
for mat, e in erows.items():
    g = lambda k: e.get(k) or 0
    pi.append(f"({q(mat)}, {g('heures Absences')!r}, {g('heures Abs Maladie')!r}, {(e.get('taux maladie') or 1)!r}, {g('nombre de panier')!r}, "
              f"{g('HS115%')!r}, {g('HS140%')!r}, {g('HS160%')!r}, {g('HS200%')!r}, {g('Congé')!r}, {g('appoint m-1')!r})")
    for col, code in addmap.items():
        v = e.get(col)
        if v:
            adds.append(f"({q(mat)}, {q(code)}, {abs(v)!r})")
d.append("""insert into public.payroll_inputs (employee_id, period, absence_hours, sick_leave_hours, sick_leave_rate, meal_count,
  overtime_15, overtime_40, overtime_60, overtime_100, leave_allowance, previous_rounding)
select e.id, '2026-08-01', v.a, v.s, v.sr, v.m, v.h1, v.h2, v.h3, v.h4, v.c, v.p
from (values
""" + ",\n".join(pi) + """
) as v(employee_number, a, s, sr, m, h1, h2, h3, h4, c, p)
join public.employees e on e.employee_number = v.employee_number;
""")
if adds:
    d.append("""insert into public.additional_salaries (employee_id, period, component_code, amount)
select e.id, '2026-08-01', v.code, v.amount
from (values
""" + ",\n".join(adds) + """
) as v(employee_number, code, amount)
join public.employees e on e.employee_number = v.employee_number;
""")
open(out_data, 'w').write("\n".join(d))

# référence Excel pour la validation
keys = {'BRUT':'gross_pay','Impôt sur le revenu':'income_tax','Retenue T.R.I.M.F':'trimf','TOTAL RETENUES':'total_employee_deductions',
        'CHARGES PATRONALES':'employer_charges','Remunération nette':'net_remuneration','NET A  PAYER':'net_pay',
        'MASSE  SALARIALE':'payroll_cost','PRIME D\'ANCIENNETE':'seniority_bonus','C.F.C.E.':'cfce',
        'I.P.R.E.S R.G part salarié':'ipres_rg_employee','I.P.R.E.S R.G. part patronale':'ipres_rg_employer',
        'CSS A.F.':'css_af','CSS A.T.':'css_at','IPM PART SALARIALE':'ipm_employee','indemnité de transport':'transport_allowance',
        'PROVISION CONGE du mois':'leave_provision'}
ref_json = {mat: {v: e.get(k) for k, v in keys.items()} for mat, e in erows.items()}
json.dump(ref_json, open('/home/claude/gen/excel_expected.json','w'), indent=1, default=str)
print(len(rows), 'employees;', len(grades), 'grades;', len(adds), 'additional salaries;', len(agreements), 'agreements')
print('masse key present:', [k for k in eh if 'MASSE' in k])
